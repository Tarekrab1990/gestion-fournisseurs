#!/usr/bin/env python3
"""Importe le classeur Excel de suivi dans la base de l'application.

    python scripts/importer_excel.py chemin/vers/suivi.xlsx [--remplacer]

Le classeur doit comporter deux feuilles : « Fournisseurs » (en-têtes ligne 5,
données à partir de la ligne 6) et « Paiements » (en-têtes ligne 7, données à
partir de la ligne 8), comme le modèle d'origine. Les colonnes calculées du
classeur (expiration du C20, alertes, statuts, contrôles) sont ignorées :
l'application les recalcule.

Les lignes en doublon (même nom de fournisseur, même n° de bon) sont ignorées,
sauf avec --remplacer qui met à jour l'enregistrement existant.
"""

import argparse
import os
import sys
from datetime import datetime, date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl est requis : pip install openpyxl")

from app import create_app
from app.database import query, execute, now

ETATS = {
    "facture originale remise": "Facture remise",
    "facture remise": "Facture remise",
    "facture originale non remise": "Facture non remise",
    "facture non remise": "Facture non remise",
}
MODES = {"cheque": "Chèque", "chèque": "Chèque", "virement": "Virement"}


def texte(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    return "" if s.startswith("#") else s          # écarte #REF!, #N/A…


def iso(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = texte(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def nombre(v):
    if v in (None, ""):
        return 0.0
    try:
        return float(str(v).replace(" ", "").replace("\u202f", "").replace(",", "."))
    except ValueError:
        return 0.0


def importer_fournisseurs(ws, remplacer):
    ajoutes = ignores = 0
    for row in ws.iter_rows(min_row=6, values_only=True):
        nom = texte(row[1] if len(row) > 1 else None)
        if not nom:
            continue
        c20 = iso(row[6] if len(row) > 6 else None)
        vals = (texte(row[2]), texte(row[3]), texte(row[4]), texte(row[5]),
                1 if c20 else 0, c20, texte(row[10] if len(row) > 10 else None),
                iso(row[11] if len(row) > 11 else None), texte(row[12] if len(row) > 12 else None))
        existe = query("SELECT id FROM suppliers WHERE name = ?", (nom,), one=True)
        if existe and not remplacer:
            ignores += 1
            continue
        if existe:
            execute("UPDATE suppliers SET department=?, nif=?, rc=?, rib=?, c20_provided=?,"
                    " c20_date=?, contact=?, last_reminder=?, note=?, updated_at=?,"
                    " updated_by='import' WHERE id=?", vals + (now(), existe["id"]))
        else:
            execute("INSERT INTO suppliers (name, department, nif, rc, rib, c20_provided,"
                    " c20_date, contact, last_reminder, note, created_at, updated_at, updated_by)"
                    " VALUES (?,?,?,?,?,?,?,?,?,?,?,?, 'import')", (nom,) + vals + (now(), now()))
        ajoutes += 1
    return ajoutes, ignores


def importer_bons(ws, remplacer):
    ajoutes = ignores = 0
    for row in ws.iter_rows(min_row=8, values_only=True):
        numero = texte(row[4] if len(row) > 4 else None)
        po_date = iso(row[3] if len(row) > 3 else None)
        if not numero or not po_date:
            continue
        etat = ETATS.get(texte(row[10]).lower(), "")
        mode = MODES.get(texte(row[11]).lower(), "")
        vals = (po_date, texte(row[2]), texte(row[5]), texte(row[6]), nombre(row[7]),
                texte(row[8]), iso(row[9]), etat, mode,
                texte(row[12]) if mode != "Virement" else "",
                texte(row[13]) if mode != "Chèque" else "",
                iso(row[14]), iso(row[15]),
                1 if texte(row[20] if len(row) > 20 else None).lower() == "oui" else 0,
                texte(row[23] if len(row) > 23 else None))
        existe = query("SELECT id FROM purchase_orders WHERE po_number = ?", (numero,), one=True)
        if existe and not remplacer:
            ignores += 1
            continue
        if existe:
            execute("UPDATE purchase_orders SET po_date=?, department=?, supplier_name=?,"
                    " description=?, amount_ttc=?, invoice_number=?, invoice_date=?,"
                    " invoice_status=?, payment_mode=?, cheque_number=?, transfer_number=?,"
                    " payment_date=?, debit_date=?, quote_done=?, note=?, updated_at=?,"
                    " updated_by='import' WHERE id=?", vals + (now(), existe["id"]))
        else:
            execute("INSERT INTO purchase_orders (po_date, department, supplier_name, description,"
                    " amount_ttc, invoice_number, invoice_date, invoice_status, payment_mode,"
                    " cheque_number, transfer_number, payment_date, debit_date, quote_done, note,"
                    " po_number, created_at, updated_at, updated_by)"
                    " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?, 'import')",
                    vals + (numero, now(), now()))
        ajoutes += 1
    return ajoutes, ignores


def main():
    ap = argparse.ArgumentParser(description="Importe un classeur de suivi dans la base.")
    ap.add_argument("fichier")
    ap.add_argument("--remplacer", action="store_true",
                    help="met à jour les enregistrements déjà présents au lieu de les ignorer")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.fichier, data_only=True)
    app = create_app()
    with app.app_context():
        total = []
        if "Fournisseurs" in wb.sheetnames:
            a, i = importer_fournisseurs(wb["Fournisseurs"], args.remplacer)
            total.append(f"fournisseurs : {a} importés, {i} ignorés")
        if "Paiements" in wb.sheetnames:
            a, i = importer_bons(wb["Paiements"], args.remplacer)
            total.append(f"bons de commande : {a} importés, {i} ignorés")
        if not total:
            sys.exit("Aucune feuille « Fournisseurs » ou « Paiements » trouvée dans le classeur.")
        print("\n".join(total))
        print("Les devis comparatifs, dates de devis et signatures doivent être vérifiés "
              "dans le module « Devis comparatifs » après import.")


if __name__ == "__main__":
    main()
